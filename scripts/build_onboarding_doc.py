"""Regenerate workflow-contract-onboarding-guide.docx.

One-off tool, not part of `make check` or the validator pipeline (those are
stdlib-only). Prerequisite:

    pip install python-docx

The committed .docx is the reproducible output. Rebuilding the rendered
PDF/PNG additionally needs LibreOffice or pandoc.
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import os

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT = os.path.join(REPO_ROOT, "workflow-contract-onboarding-guide.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, name):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)


def set_para_spacing(paragraph, before=0, after=0, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "F5F7FA")
    p = cell.paragraphs[0]
    set_para_spacing(p, before=0, after=0, line=1.15)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        run = p.add_run(line)
        set_run_font(run, "Courier New")
        run.font.size = Pt(9.5)


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    set_para_spacing(p, after=4, line=1.15)
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_para_spacing(p, after=4, line=1.15)
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    return p, run


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for level, size, before, after, color in [
    (1, 16, 18, 10, RGBColor(0x2E, 0x74, 0xB5)),
    (2, 13, 14, 7, RGBColor(0x2E, 0x74, 0xB5)),
    (3, 12, 10, 5, RGBColor(0x1F, 0x4D, 0x78)),
]:
    style = styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.15

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_para_spacing(title, after=6, line=1.0)
run = title.add_run("Workflow Contract and Skill Onboarding Guide")
run.bold = True
run.font.size = Pt(24)
run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
set_run_font(run, "Calibri")

subtitle = doc.add_paragraph()
set_para_spacing(subtitle, after=10, line=1.15)
subrun = subtitle.add_run(
    "How the workflow works, how to install it, how to publish the skill on skills.sh, and how to onboard a new operator."
)
subrun.italic = True
subrun.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

meta = doc.add_paragraph()
set_para_spacing(meta, after=12, line=1.15)
meta_run = meta.add_run("Repository: danfordChris/workflow-doc")
meta_run.font.size = Pt(10)
meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

add_heading(doc, "1. What This Package Is", 1)
doc.add_paragraph(
    "Workflow Contract is a planning-first engineering workflow for teams and agents. "
    "It separates unresolved thinking, approved design truth, execution planning, implementation, "
    "and review so that code and documentation do not drift apart."
)
doc.add_paragraph(
    "This repository contains both the canonical workflow package and a publishable companion skill, "
    "`workflow`, which helps an agent classify work correctly and generate PRD and TRD artifacts without collapsing design and implementation layers."
)

add_heading(doc, "2. Core Concepts", 1)
add_bullet(doc, "Thinking mode: research, decisions, design truth, proposals, and planning. No code changes.")
add_bullet(doc, "Execution mode: implementation against a prepared task with explicit scope and acceptance criteria.")
add_bullet(doc, "Review mode: reconcile delivered work against the design truth and planning artifacts.")
add_bullet(doc, "Design layer: approved requirements, product behavior, contracts, and architecture in `docs/design/`.")
add_bullet(doc, "Implementation layer: project plans, tasks, sequencing, rollout, and status in `docs/implementation/`.")
add_bullet(doc, "Changes layers: use `docs/changes/wayfinding/` for large unclear work and `docs/changes/proposed/` for concrete but unresolved deltas.")

add_heading(doc, "3. PRD And TRD Mapping", 1)
doc.add_paragraph(
    "In this workflow, a PRD and a TRD are separate artifacts with different authority. "
    "They should not be combined into one mixed document."
)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.autofit = False
widths = [Inches(1.3), Inches(2.0), Inches(3.2)]
for idx, width in enumerate(widths):
    table.columns[idx].width = width
hdr = table.rows[0].cells
for i, text in enumerate(["Artifact", "Primary Location", "Purpose"]):
    hdr[i].text = text
    set_cell_shading(hdr[i], "E8EEF5")
    hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
for text in [
    ("PRD", "`docs/design/<feature>.md`", "Defines user problem, product behavior, requirements, constraints, and acceptance truth."),
    ("TRD", "`docs/implementation/<feature>.md`", "Defines technical approach, delivery slices, dependencies, rollout, verification, and task planning."),
]:
    cells = table.add_row().cells
    for i, value in enumerate(text):
        cells[i].text = value
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

add_heading(doc, "4. What Is In This Repository", 1)
add_bullet(doc, "`spec/`: canonical workflow policy, lifecycle, guardrails, and task standards.")
add_bullet(doc, "`scripts/validate_workflow.py`: validator for structure, metadata, readiness, scope conflicts, transitions, and references.")
add_bullet(doc, "`templates/`: reusable workflow document templates.")
add_bullet(doc, "`docs/`: example and live documentation scaffold owned by the workflow.")
add_bullet(doc, "`.agents/skills/workflow/`: publishable companion skill for operator guidance and PRD/TRD generation.")
add_bullet(doc, "`AGENTS.md`: repository startup instructions and workflow authority declaration.")

add_heading(doc, "5. Installation Paths", 1)
doc.add_paragraph(
    "There are two common installation paths. One installs the workflow package into another repository. "
    "The other installs the companion skill from skills.sh-compatible GitHub distribution."
)

add_heading(doc, "5.1 Install The Workflow In A Consumer Repository", 2)
add_number(doc, "From the target repository root, add the workflow contract as a submodule.")
add_code_block(
    doc,
    [
        "git submodule add git@github.com:danfordChris/workflow-doc.git .agents/workflows/workflow-contract",
        "git submodule update --init --recursive",
    ],
)
add_number(doc, "Review or create `AGENTS.md` and ensure it declares workflow authority and startup order.")
add_number(doc, "Bootstrap and validate the installed workflow.")
add_code_block(
    doc,
    [
        "make -C .agents/workflows/workflow-contract check",
        "python3 .agents/workflows/workflow-contract/scripts/validate_workflow.py",
    ],
)
add_number(doc, "Confirm that `docs/design/`, `docs/implementation/`, and the changes folders are present and match the workflow contract.")

add_heading(doc, "5.2 Manual Fallback (No Make)", 2)
add_code_block(
    doc,
    [
        "python3 .agents/workflows/workflow-contract/scripts/init_workflow_contract.py",
        "python3 .agents/workflows/workflow-contract/scripts/validate_workflow.py",
    ],
)

add_heading(doc, "5.3 Install The Companion Skill", 2)
doc.add_paragraph(
    "Install the repository as a skills source when you want agent operators to pull the companion skill through the `skills` CLI."
)
add_code_block(
    doc,
    [
        "npx skills add danfordChris/workflow-doc",
        "npx skills add danfordChris/workflow-doc --skill workflow",
    ],
)
add_bullet(doc, "The install flow is interactive and asks which agents to install to, the installation scope, and the installation method.")
add_bullet(doc, "A symlink-based project installation is typically the simplest local setup.")
add_bullet(doc, "The repository was successfully discovered by `skills` and exposes `workflow` as an installable skill.")

add_heading(doc, "6. How The Skill Works", 1)
doc.add_paragraph(
    "The `workflow` skill is not the canonical workflow policy. "
    "It is an operator playbook that routes agents through the repo's actual source of truth."
)
add_bullet(doc, "First it classifies the request by layer, lifecycle state, and agent mode.")
add_bullet(doc, "Then it forces the agent to read the canonical spec files in order before changing workflow artifacts.")
add_bullet(doc, "It applies layer rules so that design truth does not leak into planning documents and task lists do not invent product behavior.")
add_bullet(doc, "When asked for a PRD, it places the document in `docs/design/`.")
add_bullet(doc, "When asked for a TRD, it places the document in `docs/implementation/`.")
add_bullet(doc, "When both are needed, it generates the PRD first and derives the TRD from that approved design input.")

add_heading(doc, "7. First-Day Onboarding Checklist", 1)
for item in [
    "Read `README.md` for the high-level operating model.",
    "Read `spec/workflow-spec.md`, `spec/guardrails-spec.md`, `spec/lifecycle-spec.md`, and `spec/task-spec.md` in that order.",
    "Read `AGENTS.md` to understand repo-specific startup rules.",
    "Run the validator once so you know what a passing repo looks like.",
    "Inspect the companion skill at `.agents/skills/workflow/SKILL.md`.",
    "Install the skill locally if you will operate through the `skills` CLI.",
    "Create one sample PRD and one sample TRD so the layer separation becomes concrete.",
]:
    add_number(doc, item)

add_heading(doc, "8. Recommended Daily Workflow", 1)
add_number(doc, "Classify the work: design, implementation, wayfinding, or proposal.")
add_number(doc, "If the route is unclear, start in `docs/changes/wayfinding/` before writing a proposal or PRD.")
add_number(doc, "If the behavior is concrete but unresolved, capture it in `docs/changes/proposed/`.")
add_number(doc, "Move approved behavior into `docs/design/`.")
add_number(doc, "Turn approved truth into implementation planning in `docs/implementation/`.")
add_number(doc, "Execute one bounded task at a time.")
add_number(doc, "Review delivered work against both the code and the documented truth.")

add_heading(doc, "9. Publishing On skills.sh", 1)
doc.add_paragraph(
    "skills.sh does not use a manual submit form in the normal install flow. "
    "The repository becomes discoverable through public GitHub availability plus installs via the `skills` CLI."
)
add_bullet(doc, "Keep the GitHub repository public.")
add_bullet(doc, "Keep a valid `SKILL.md` with frontmatter in the skill folder.")
add_bullet(doc, "Optionally include `skills.sh.json` to curate groupings and presentation.")
add_bullet(doc, "Trigger discovery with at least one `npx skills add danfordChris/workflow-doc` install.")
add_bullet(doc, "Wait for telemetry ingestion and cache refresh on skills.sh.")

add_heading(doc, "10. Troubleshooting", 1)
add_heading(doc, "10.1 Validator Fails", 2)
add_bullet(doc, "If `.agents/workflows/workflow-contract/` is missing, the validator will fail on required structure checks.")
add_bullet(doc, "If a task is marked in progress without scope, acceptance criteria, or session budget, readiness validation will fail.")
add_bullet(doc, "If two active tasks claim the same scope, the scope validator will fail.")

add_heading(doc, "10.2 skills CLI Install Issues", 2)
add_bullet(doc, "If the CLI repeats prompt lines, that is often just interactive redraw noise rather than a hard failure.")
add_bullet(doc, "If the skill is found and the installer reaches agent selection, repository discovery is already working.")
add_bullet(doc, "If the repository cannot be found, verify the public GitHub slug and the `SKILL.md` location.")

add_heading(doc, "10.3 PRD And TRD Drift", 2)
add_bullet(doc, "If a TRD starts defining product behavior, move that behavior back into a PRD or an accepted proposal.")
add_bullet(doc, "If a PRD contains sprint or task management, move that content into `docs/implementation/`.")

add_heading(doc, "11. Command Reference", 1)
cmd_table = doc.add_table(rows=1, cols=2)
cmd_table.alignment = WD_TABLE_ALIGNMENT.LEFT
cmd_table.autofit = False
cmd_table.columns[0].width = Inches(2.7)
cmd_table.columns[1].width = Inches(3.6)
row = cmd_table.rows[0].cells
row[0].text = "Command"
row[1].text = "Purpose"
set_cell_shading(row[0], "E8EEF5")
set_cell_shading(row[1], "E8EEF5")
commands = [
    ("git submodule add git@github.com:danfordChris/workflow-doc.git .agents/workflows/workflow-contract", "Install the workflow package into a repository."),
    ("make -C .agents/workflows/workflow-contract check", "Bootstrap and validate the installed workflow."),
    ("python3 .agents/workflows/workflow-contract/scripts/validate_workflow.py", "Run the canonical validator directly."),
    ("npx skills add danfordChris/workflow-doc", "Install the repository as a skills source."),
    ("npx skills add danfordChris/workflow-doc --skill workflow", "Install only the PRD/TRD-capable companion skill."),
]
for command, purpose in commands:
    cells = cmd_table.add_row().cells
    cells[0].text = command
    cells[1].text = purpose
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

add_heading(doc, "12. Operator Quick Start", 1)
doc.add_paragraph(
    "Use this script when onboarding a new teammate or agent operator. "
    "It keeps the first session concrete and short."
)
add_code_block(
    doc,
    [
        "1. Read AGENTS.md and spec/workflow-spec.md.",
        "2. Run the validator once.",
        "3. Install the workflow skill if you use the skills CLI.",
        "4. Draft a PRD in docs/design/ for one small feature.",
        "5. Draft the matching TRD in docs/implementation/.",
        "6. Ask the reviewer to check that design truth and implementation planning stayed separate.",
    ],
)

doc.save(OUTPUT)
print(OUTPUT)
